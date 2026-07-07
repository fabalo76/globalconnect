from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path(__file__).resolve().parent
DOCX_PATH = OUT_DIR / "NEXGO_Android_Smart_PinPAD_Programmers_Manual.docx"

FONT = "Times New Roman"
MONO = "Courier New"
BLUE = "1F4E79"
LIGHT = "EAF2F8"
GRAY = "F4F4F4"


@dataclass
class Elem:
    field: str
    length: str
    description: str


@dataclass
class Flow:
    host: str
    direction: str
    device: str


@dataclass
class Message:
    code: str
    title: str
    frame: str
    request_format: str
    response_format: str
    length: str
    usage: str
    request: list[Elem]
    response: list[Elem]
    flow: list[Flow]
    notes: list[str] = field(default_factory=list)


def para(doc: Document, text: str = "", style: str | None = None, *,
         size: float = 10.5, bold: bool = False, italic: bool = False,
         font: str = FONT, color: str = "000000", before: float = 0,
         after: float = 4, align: int | None = None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = RGBColor.from_string(color)
    return p


def run(p, text: str, *, size=10.5, bold=False, italic=False, font=FONT, color="000000"):
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = RGBColor.from_string(color)
    return r


def set_cell_text(cell, text: str, *, bold=False, font=FONT, size=9.5, align=None):
    cell.text = ""
    for idx, line in enumerate(str(text).split("\n")):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.06
        if align is not None:
            p.alignment = align
        r = p.add_run(line)
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(cell)


def set_cell_margins(cell, top=65, start=100, bottom=65, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, val in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def inches_to_dxa(value: float) -> int:
    return int(round(value * 1440))


def set_cell_width(cell, width_in: float):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(inches_to_dxa(width_in)))
    tc_w.set(qn("w:type"), "dxa")
    cell.width = Inches(width_in)


def set_fixed_table_widths(table, widths_in: list[float]):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(inches_to_dxa(w) for w in widths_in)))
    tbl_w.set(qn("w:type"), "dxa")

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_in[idx] if idx < len(widths_in) else widths_in[-1]
            set_cell_width(cell, width)


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table, color="000000", size="4"):
    tbl_pr = table._tbl.tblPr
    b = tbl_pr.first_child_found_in("w:tblBorders")
    if b is None:
        b = OxmlElement("w:tblBorders")
        tbl_pr.append(b)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        e = b.find(qn(f"w:{edge}"))
        if e is None:
            e = OxmlElement(f"w:{edge}")
            b.append(e)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), size)
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    tr_pr.append(h)


def element_table(doc: Document, caption: str, subtitle: str, rows: Iterable[Elem]):
    para(doc, caption, size=10.5, after=2)
    widths = [1.35, 0.55, 5.05]
    table = doc.add_table(rows=2, cols=3)
    set_fixed_table_widths(table, widths)
    borders(table)
    repeat_header(table.rows[1])
    hdr = table.rows[0].cells
    hdr[0].merge(hdr[2])
    set_cell_text(hdr[0], subtitle, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade(hdr[0], LIGHT)
    for idx, label in enumerate(["Field", "Length", "Value and description"]):
        c = table.rows[1].cells[idx]
        set_cell_text(c, label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for item in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], item.field, font=MONO, size=9)
        set_cell_text(cells[1], item.length, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[2], item.description)
    set_fixed_table_widths(table, widths)
    para(doc, "", after=3)


def flow_table(doc: Document, rows: Iterable[Flow]):
    para(doc, "Message flow:", size=10.5, after=2)
    widths = [2.35, 0.42, 4.18]
    table = doc.add_table(rows=1, cols=3)
    set_fixed_table_widths(table, widths)
    borders(table)
    repeat_header(table.rows[0])
    for idx, label in enumerate(["HOST", "Dir.", "PINPAD"]):
        cell = table.rows[0].cells[idx]
        set_cell_text(cell, label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(cell, LIGHT)
    for row in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], row.host)
        set_cell_text(cells[1], row.direction, align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
        set_cell_text(cells[2], row.device)
    set_fixed_table_widths(table, widths)
    para(doc, "", after=4)


def add_format(doc: Document, label: str, text: str):
    p = para(doc, "", after=0)
    run(p, label, size=10.5)
    run(p, text, size=9.5, font=MONO)


def message_section(doc: Document, msg: Message):
    para(doc, f"Message {msg.code}", size=11.5, bold=True, after=0)
    p = doc.paragraphs[-1]
    run(p, f"    {msg.title}", size=11.5, bold=True)
    add_format(doc, "Format:        ", msg.request_format)
    add_format(doc, "               ", msg.response_format)
    para(doc, "", after=6)
    p = para(doc, "", after=2)
    run(p, "Message length:  ", size=10.5)
    run(p, msg.length, size=10.5)
    p = para(doc, "", after=2)
    run(p, "Usage:           ", size=10.5)
    run(p, msg.usage, size=10.5)
    if msg.notes:
        p = para(doc, "", after=2)
        run(p, "Implementation:  ", size=10.5)
        run(p, msg.notes[0], size=10.5)
        for note in msg.notes[1:]:
            para(doc, note, size=10.5, after=2)
    element_table(doc, "Message element:", f"Request frame (HOST to PINPAD)", msg.request)
    element_table(doc, "Message element:", f"Response frame (PINPAD to HOST)", msg.response)
    flow_table(doc, msg.flow)


def req_common(command: str, frame: str, fields: list[Elem]) -> list[Elem]:
    start = "<SI>" if frame == "admin" else "<STX>"
    end = "<SO>" if frame == "admin" else "<ETX>"
    start_hex = "<0F>" if frame == "admin" else "<02>"
    end_hex = "<0E>" if frame == "admin" else "<03>"
    return [Elem(start, "1", start_hex), Elem(command, str(len(command)), "Message ID")] + fields + [
        Elem(end, "1", end_hex),
        Elem("[LRC]", "1", "Checksum"),
    ]


def resp_common(command: str, frame: str, fields: list[Elem]) -> list[Elem]:
    return req_common(command, frame, fields)


def standard_flow(command: str, response: str, final_eot: bool = False, async_response: bool = False) -> list[Flow]:
    processing = "Validate command format and process request."
    if async_response:
        processing = "Accept valid request and start the operation. Send the response when the operation completes."
    rows = [
        Flow(f"Message {command} request frame", "->", ""),
        Flow("", "<-", "<ACK>/<NAK>/<EOT>"),
        Flow("", "<-", processing + f"\nIf successful, send Message {response}. If format is invalid, send <EOT> or command-specific failure."),
        Flow("Verify response frame.\nIf valid, send <ACK>.\nIf LRC error, send <NAK>.\nIf canceling the session, send <EOT>.", "->", ""),
    ]
    if final_eot:
        rows.append(Flow("", "<-", "Send final <EOT> after host <ACK>."))
    return rows


def no_frame_flow(command: str) -> list[Flow]:
    return [
        Flow(f"Message {command} request frame", "->", ""),
        Flow("", "<-", "<ACK>/<NAK>/<EOT>"),
        Flow("", "<-", "Apply command action. No data frame is sent for a normal accepted command."),
    ]


def messages() -> list[Message]:
    admin = "admin"
    tran = "transaction"
    items: list[Message] = []

    def add(code, title, frame, req_fields, resp_fields, length, usage, response=None, notes=None, async_response=False, final_eot=None):
        if response is None:
            response = code
        if final_eot is None:
            final_eot = frame == admin and code not in {"11", "14"}
        start = "<SI>" if frame == admin else "<STX>"
        end = "<SO>" if frame == admin else "<ETX>"
        items.append(Message(
            code=code,
            title=title,
            frame=frame,
            request_format=f"{start}{code}[Data]{end}[LRC] (request frame)",
            response_format=f"{start if frame == admin else '<STX>'}{response}[Data]{end if frame == admin else '<ETX>'}[LRC] (response frame)",
            length=length,
            usage=usage,
            request=req_common(code, frame, req_fields),
            response=resp_common(response, frame, resp_fields),
            flow=standard_flow(code, response, final_eot=final_eot, async_response=async_response),
            notes=notes or [],
        ))

    def add_no_frame(code, title, frame, req_fields, length, usage, notes=None):
        start = "<SI>" if frame == admin else "<STX>"
        end = "<SO>" if frame == admin else "<ETX>"
        items.append(Message(
            code=code,
            title=title,
            frame=frame,
            request_format=f"{start}{code}[Data]{end}[LRC] (request frame)",
            response_format="<ACK>/<NAK>/<EOT> only unless an error frame is specified.",
            length=length,
            usage=usage,
            request=req_common(code, frame, req_fields),
            response=[Elem("<ACK>", "1", "Command accepted"), Elem("<NAK>", "1", "Transmission error"), Elem("<EOT>", "1", "Format error, timeout, or canceled session")],
            flow=no_frame_flow(code),
            notes=notes or [],
        ))

    add("02", "Load Master Key or Working Key", admin,
        [Elem("[Key ID]", "1", "'0' to '9' and 'A' to 'G'."), Elem("[Key value]", "16, 32, or 48", "Hexadecimal key value. Single, double, and triple length keys are accepted according to the selected algorithm."), Elem("<FS>", "1", "Optional field separator. Present when key attributes follow."), Elem("[Usage]", "2", "Optional key usage. K0 key encryption, K1 key transport, P0 PIN encryption, M3 MAC, D0 data encryption."), Elem("[Mode]", "1", "Optional key mode. Common values include D, E, G, V."), Elem("[Algorithm]", "1", "Optional algorithm. Default is T when omitted.")],
        [Elem("02", "2", "Message ID echoed with the original payload on success."), Elem("?", "1", "Present only in the error frame."), Elem("[Error code]", "1", "Command-specific key loading error code.")],
        "Variable.", "Loads a clear key into the selected secure key slot or working-key role.",
        notes=["K0 writes the device master-key slot. P0, M3, and D0 are stored as working PIN, MAC, and data keys. P0 uses a random wrapper master key internally."])
    add("04", "Check Master Key", admin,
        [Elem("[Key ID]", "1", "Key slot to query."), Elem("[Attribute flag]", "1", "Optional. '1' requests key attribute information.")],
        [Elem("[Status]", "1", "'F' key exists, '0' key does not exist."), Elem("[Usage]", "2", "Returned only when requested and available."), Elem("<FS>", "1", "Field separator."), Elem("[Mode]", "1", "Returned key mode."), Elem("[Algorithm]", "1", "Returned algorithm.")],
        "Variable.", "Queries whether a key exists and optionally returns the stored key usage metadata.",
        notes=["If secure key material exists but local metadata was removed by reinstall, PIN commands fail safely and require key injection again."])
    add("05", "Load Serial Number", admin, [Elem("[Serial number]", "0 to 16", "Compatibility serial number value.")], [Elem("[Serial number]", "0 to 16", "Echo of accepted serial substring.")], "Variable.", "Compatibility command. The hardware serial number is not changed.")
    add("06", "Get Serial Number", admin, [], [Elem("[Serial number]", "Variable", "Device serial number from Android or NEXGO SDK.")], "Fixed request.", "Returns the device serial number.")
    add("08", "Select Active Master Key", admin, [Elem("[Key ID]", "1", "Key id to select.")], [Elem("[Status]", "1", "'0' success, '1' failure.")], "Fixed.", "Selects the active key used by MK/SK PIN and MAC commands.")
    add("09", "Communications Test", admin, [], [Elem("<SUB>", "1", "Subfield delimiter."), Elem("PROCESSING", "10", "Processing echo text, followed by final status in the second exchange.")], "Fixed.", "Performs a two-step communication loopback test.", notes=["The host must echo the first response payload. The device then returns status 0 for success."])
    add_no_frame("11", "Device Connection Test", admin, [], "Fixed.", "Connection probe. A valid frame is acknowledged without a final EOT.")
    add("12", "Select Prompt Language", admin, [Elem("[Language index]", "1", "Language table index requested by host.")], [Elem("[Status]", "1", "'0' accepted, '1' rejected.")], "Fixed.", "Compatibility prompt-language command.", notes=["The Android application language follows Android locale settings."])
    add("13", "Set Communication Speed", admin, [Elem("[Baud code]", "1", "Baud-rate selector 1 through 8."), Elem("[Mode]", "0 or 1", "Optional communication mode selector.")], [Elem("[Status]", "1", "'0' accepted, '1' rejected.")], "Fixed.", "Validates a communication speed request.")
    add_no_frame("14", "Acknowledge Without Final End", admin, [], "Fixed.", "Compatibility acknowledgement command.")
    add("17", "Request Random Number", admin, [], [Elem("[Random]", "16", "Eight random bytes encoded as hexadecimal.")], "Fixed.", "Returns an 8-byte random challenge.")
    add("19", "Query Firmware Version", admin, [Elem("[Component]", "1", "Firmware component selector."), Elem("[Option]", "0 or 1", "Optional component option.")], [Elem(".", "1", "Version prefix separator."), Elem("[Version]", "Variable", "Application, SDK, or device firmware version string.")], "Variable.", "Returns component firmware or application version.", notes=["The version string is formatted for compatibility and includes the NEXGO model when available."])
    add("1C", "Query Hardware Capability", admin, [], [Elem("<FS>", "1", "Field separator before each capability item."), Elem("[Capability]", "Variable", "Capability string.")], "Variable.", "Returns supported hardware and application capability values.")
    add("1F", "Query Prompt Table", admin, [], [Elem("[Language list]", "Variable", "Prompt-language identifiers separated by <FS>.")], "Variable.", "Reports usable prompt tables.")
    add("1M", "Setup Keypad Beeper", admin, [Elem("[Option]", "1", "'0' disable, '1' enable.")], [Elem("[Status]", "1", "Echoed option on success. '0' on invalid request.")], "Fixed.", "Controls keypad beep preference.")
    add("1P", "Control Beeper", admin, [Elem("<SUB>", "1", "Subfield delimiter."), Elem("[Count]", "1", "Number of beeps, 1 through 9."), Elem("<FS>", "1", "Field separator."), Elem("[Duration]", "1 to 3", "Duration in 10 ms units."), Elem("[Interval]", "1 to 3", "Interval in 10 ms units.")], [Elem("[Status]", "1", "'0' success, '1' failure.")], "Variable.", "Plays one or more beeps.")
    add("M03", "Load Permanent Unit Serial Number", admin, [Elem("[Serial]", "11", "Uppercase alphanumeric or dash serial value."), Elem("[Slot]", "0 or 1", "Optional slot 0 through 7. Default 0.")], [Elem("[Status]", "1", "'0' success, '1' bad serial, '2' already set, '3' storage or slot error.")], "Variable.", "Stores a reinstall-persistent private serial number.")
    add("M04", "Query Permanent Unit Serial Number", admin, [Elem("[Slot]", "0 or 1", "Optional slot. Default 0.")], [Elem("[Serial]", "0 or 11", "Stored permanent serial, empty if none was loaded.")], "Variable.", "Returns the permanent unit serial number.")

    add("20", "Load Secret Master Key", tran, [Elem("[Option]", "1", "'0' or '1'."), Elem("[Key value]", "16 or 32", "Secret master key in hexadecimal.")], [Elem("20", "2", "Message ID echoed with original payload on success.")], "Variable.", "Loads the secret master key used by secret PIN session slots.", notes=["The key is injected into internal secure slots 20 through 29. Session-key loading waits until this master key is ready."], final_eot=False)
    add("21", "Load Secret Session Key", tran, [Elem("[Key ID]", "1", "Secret session key id 0 through 9."), Elem("[Key value]", "16 or 32", "Encrypted PIN session key in hexadecimal.")], [Elem("21", "2", "Message ID echoed with original payload on success.")], "Variable.", "Loads a secret MK/SK session key for PIN entry.", notes=["Key id 0 through 9 maps internally to secure slots 20 through 29."], final_eot=False)
    for code, title in [("22", "PIN Entry Request Using Secret MK/SK"), ("23", "PIN Entry Request With External Prompt"), ("24", "PIN Entry Request With Custom Prompt")]:
        add(code, title, tran, [Elem("[Account]", "8 to 19", "PAN/account digits used for PIN block formatting."), Elem("<FS>", "1", "Field separator."), Elem("[Key/control]", "Variable", "Secret key id followed by optional PIN control and prompt data.")], [Elem("[Status/PIN data]", "Variable", "Message 71 payload containing PIN block or error code.")], "Variable.", "Starts secure online PIN entry using the secret key family.", response="71", async_response=True, final_eot=False)
    for code, title, usage in [("60", "DUKPT PIN Entry Request", "Starts online PIN entry using DUKPT."), ("70", "PIN Entry Request", "Starts DUKPT or MK/SK PIN entry. A leading dot selects MK/SK."), ("7G", "PIN Change Request", "Starts MK/SK PIN change style entry."), ("Z60", "External Prompt PIN Entry", "Starts PIN entry using an external prompt."), ("Z62", "Custom Prompt PIN Entry", "Starts PIN entry with custom prompt lines and PIN limits.")]:
        add(code, title, tran, [Elem("[Account]", "8 to 19", "PAN/account digits."), Elem("<FS>", "1", "Field separator."), Elem("[Session key/control]", "Variable", "Session key, PIN length control, optional prompts, and optional timeout.")], [Elem("[PIN response]", "Variable", "Message 71 payload. Includes status, PIN block, KSN/key id, or error.")], "Variable.", usage, response="71", async_response=True, final_eot=False)
    add("62", "Amount Authorization", tran, [Elem("[Data]", "Variable", "Compatibility payload.")], [Elem("[Status]", "1", "'0' accepted.")], "Variable.", "Compatibility command that returns message 63.", response="63")
    for code in ["63", "71", "91", "99", "Z65", "Z67"]:
        add_no_frame(code, "Response Message Compatibility Acknowledgement", tran, [Elem("[Payload]", "Variable", "Ignored when received from host.")], "Variable.", "Accepted as a no-op when a host sends a response-id command.")
    add_no_frame("7A", "Set DUKPT KSN Output Format", tran, [Elem("[Format]", "1", "KSN output format selector.")], "Fixed.", "Controls DUKPT KSN formatting.")
    add("90", "Load DUKPT Initial Key Set 0", tran, [Elem("[IPEK]", "32", "DUKPT initial key."), Elem("[KSN]", "20", "Initial key serial number.")], [Elem("[Status]", "1", "'0' success. '1' followed by reason on failure.")], "Fixed.", "Injects the first DUKPT key set.", response="91")
    add("94", "Load DUKPT Initial Key Set 1", tran, [Elem("[IPEK]", "32", "DUKPT initial key."), Elem("[KSN]", "20", "Initial key serial number.")], [Elem("[Status]", "1", "'0' success. '1' followed by reason on failure.")], "Fixed.", "Injects the second DUKPT key set.", response="91")
    add_no_frame("96", "Select DUKPT Key Set", tran, [Elem("[Key set]", "1", "'0' or '1'.")], "Fixed.", "Selects active DUKPT key set.")
    add("98", "Query DUKPT Key Set", tran, [Elem("[Key set]", "1", "'0' or '1'."), Elem("[Info flag]", "0 or 1", "Optional '1' requests extra key information.")], [Elem("[Status]", "1", "Key state."), Elem("[Info]", "Variable", "Optional algorithm/type/KCV style fields separated by <FS>.")], "Variable.", "Queries DUKPT key state.", response="99")
    add("Z64", "Query Master Key KCV", tran, [Elem("[Key ID]", "1", "Master key id.")], [Elem("[Key ID]", "1", "Echoed key id."), Elem("[KCV]", "6", "Three-byte key check value, or '?' when unavailable.")], "Fixed.", "Calculates the master key check value.", response="Z65")
    add("Z66", "Generate MAC", tran, [Elem("[MAC data]", "Variable", "Data to MAC according to active key metadata.")], [Elem("[Status]", "1", "'0' success, '2' failure."), Elem("[MAC]", "16", "MAC value when successful.")], "Variable.", "Generates a MAC using the active MAC-capable key.", response="Z67")

    for code, title, usage, response in [
        ("Q1", "Arm Magnetic Stripe Reader", "Displays the swipe prompt and waits for a swipe. Cached multi-interface swipe data is returned immediately when available.", "81"),
        ("Q8", "Read Contactless Track Data", "Starts RF card read and returns track-equivalent data.", "83"),
        ("Q9", "Read Swipe or Contactless Track Data", "Arms both swipe and RF track paths. The first successful reader wins.", "81/83"),
        ("QF", "Arm Magnetic Stripe Reader With Image Prompt", "Compatibility image-prompt variant. Android prompt animation is used.", "81"),
        ("QG", "Read Contactless Track Data With Image Prompt", "Compatibility image-prompt variant. Android prompt animation is used.", "83"),
        ("QH", "Read Swipe or Contactless Track Data With Image Prompt", "Compatibility dual-reader image-prompt variant.", "81/83"),
    ]:
        add(code, title, tran, [], [Elem("[Track 1]", "Variable", "Track 1 data when present."), Elem("<FS>", "1", "Field separator."), Elem("[Track 2]", "Variable", "Track 2 data when present."), Elem("[Track 3]", "Variable", "Track 3 data when present.")], "Variable.", usage, response=response, async_response=True, final_eot=False)
    for code, title, usage in [("Q2", "Transaction Complete", "Completes a magnetic stripe/card-read session and returns to idle."), ("Q3", "Ignore Card Swipe", "Stops active swipe handling and ignores the current swipe."), ("Q4", "Set Swipe Track Mode", "Sets swipe track mode: 0 track 2 only, 1 disabled, 2 all tracks."), ("Q5", "Set Swipe Retry Count", "Sets retry count 0 through 9."), ("QA", "Set Contactless Track Mode", "Sets contactless track mode: 0 track 2 only, 1 disabled, 2 all tracks.")]:
        add_no_frame(code, title, tran, [Elem("[Option]", "Variable", "Command-specific option.")], "Variable.", usage)
    add("Q6", "Set Swipe Output Format", tran, [Elem("[Format]", "1", "'0' clear text, '2' clear text with sentinels.")], [Elem("Q6", "2", "Request frame echoed on success.")], "Fixed.", "Sets magnetic stripe output format.")
    add("Q7", "Query Swipe Output Format", tran, [], [Elem("[Format]", "1", "Current output format.")], "Fixed.", "Returns magnetic stripe output format.")
    add("QB", "Set Swipe Auto Arm", tran, [Elem("[Option]", "1", "'0' disable, '1' enable.")], [Elem("[Status]", "1", "'0' success, '1' failure.")], "Fixed.", "Enables or disables automatic swipe arming.")
    add("QC", "Set Contactless Auto Arm", tran, [Elem("[Option]", "1", "'0' disable, '1' enable.")], [Elem("[Status]", "1", "'0' success, '1' failure.")], "Fixed.", "Enables or disables contactless auto-arm compatibility flag.")
    add("QD", "Set Contactless Output Format", tran, [Elem("[Format]", "1", "'0' clear text, '2' clear text with sentinels.")], [Elem("QD", "2", "Request frame echoed on success.")], "Fixed.", "Sets contactless track-equivalent output format.")
    add("QI", "Query Contactless Tags", tran, [Elem("[DOL]", "Variable", "Tag list or data object list. Optional ready-control byte may precede it.")], [Elem("[Status/control]", "Variable", "'F' for no data or '11' plus TLV data."), Elem("<FS>", "1", "Field separator before TLV data.")], "Variable.", "Returns requested contactless transaction data objects.", response="QJ")
    add("QK", "Multi-Interface Card Detection", tran, [], [Elem("[Result]", "1", "'1' swipe, '2' chip, '3' contactless, '0' error/cancel."), Elem("[Error]", "0 or 1", "Error code when result is 0.")], "Fixed.", "Prompts for swipe, insert, or tap and returns the detected interface.", async_response=True, final_eot=False)
    add("81", "Magnetic Stripe Data", tran, [], [Elem("[Track data]", "Variable", "Track 1, track 2, and track 3 separated by <FS>.")], "Variable.", "Device-to-host track-data response.", final_eot=False)
    add("83", "Contactless Track Data", tran, [], [Elem("[Track data]", "Variable", "Track-equivalent data separated by <FS>.")], "Variable.", "Device-to-host RF track-equivalent response.", final_eot=False)

    cpu_defs = [
        ("I00", "Query Primary Smart Card Presence", [], [Elem("[Status]", "1", "'F' present, '0' absent.")], "Checks ICC card presence.", "I00"),
        ("I01", "Primary Smart Card Cold Reset", [], [Elem("[ATR]", "Variable", "ATR returned in message I02.")], "Powers on ICC card and returns ATR.", "I02"),
        ("I04", "Primary Smart Card Deactivate", [], [], "Powers off ICC card.", "I04"),
        ("I06", "Primary Smart Card APDU Exchange", [Elem("[APDU]", "8 to 524", "APDU command in hexadecimal.")], [Elem("[APDU response]", "Variable", "Response APDU returned in message I07.")], "Exchanges an APDU with the powered ICC card.", "I07"),
        ("I11", "SAM Cold Reset", [], [Elem("[ATR]", "Variable", "ATR returned in message I12.")], "Powers on selected SAM and returns ATR.", "I12"),
        ("I14", "SAM Deactivate", [], [], "Powers off selected SAM.", "I14"),
        ("I15", "Select SAM Interface", [Elem("[Interface]", "1", "'1' PSAM1, '2' PSAM2.")], [Elem("[Interface]", "1", "Selected interface echoed.")], "Selects active SAM slot.", "I15"),
        ("I16", "SAM APDU Exchange", [Elem("[APDU]", "8 to 524", "APDU command in hexadecimal.")], [Elem("[APDU response]", "Variable", "Response APDU returned in message I17.")], "Exchanges an APDU with the powered SAM.", "I17"),
    ]
    for code, title, req, resp, usage, response in cpu_defs:
        add(code, title, tran, req, resp or [Elem("[Status]", "0", "No data on success.")], "Variable.", usage, response=response)
    for code in ["I02", "I03", "I05", "I07", "I08", "I09", "I0A", "I0B", "I0C", "I0D", "I0E", "I0F", "I12", "I17"]:
        add(code, "CPU Card Compatibility or Response Identifier", tran, [Elem("[Payload]", "Variable", "Ignored or unsupported when received from host.")], [Elem("[Error]", "2", "I0F error code for unsupported host request.")], "Variable.", "Reserved, unsupported, or device response identifier.", response="I0F")

    mifare = [
        ("P01", "Enable or Disable RF Memory Card Reader", "[0/1]", "Opens or closes the RF reader."),
        ("P02", "Query RF Card Presence", "", "Returns ATQA when a Type A card is present."),
        ("P03", "RF Anticollision", "", "Returns UID."),
        ("P04", "RF Select", "", "Returns SAK."),
        ("P05", "RF Activate", "", "Returns ATQA, SAK, and UID."),
        ("P06", "RF Halt", "", "Halts or deselects the current card."),
        ("P07", "Authenticate Classic Sector", "[Sector][Key data]", "Authenticates a Classic sector using direct or stored key."),
        ("P08", "Read Ultralight Page", "[Page]", "Reads a 4-byte Ultralight page."),
        ("P09", "Write Ultralight Page", "[Page][Data]", "Writes a 4-byte Ultralight page."),
        ("P10", "Read Block", "[Block]", "Reads a block or page."),
        ("P11", "Write Block", "[Block][Data]", "Writes a block or page."),
        ("P12", "Read Sector", "[Sector]", "Reads all blocks in a sector."),
        ("P13", "Write Sector", "[Sector][Data]", "Writes sequential blocks in a sector."),
        ("P14", "Value Operation", "[Block][Mode][Value/Target]", "Performs value write, increment, decrement, or backup operation."),
        ("P15", "Load RF Key", "[KeyNo][KeyA][KeyB]", "Stores key A and key B for later authentication."),
        ("P16", "Identify RF Card Type", "", "Returns card type code."),
        ("P17", "Activate DESFire", "", "Activates DESFire and returns ATS or UID."),
        ("P18", "Deselect DESFire", "", "Deselects DESFire and clears RF state."),
        ("P19", "RF APDU Exchange", "[APDU]", "Exchanges APDU or raw command depending on detected RF type."),
        ("P20", "RF Raw Block Exchange", "[CRC][Wait][Data]", "Performs raw RF command exchange."),
    ]
    for code, title, payload, usage in mifare:
        add(code, title, tran, [Elem("[Data]", "Variable", payload or "No payload.")], [Elem("[Status]", "1", "'0' success, '1' failure."), Elem("[Data]", "Variable", "Command-specific response data when successful.")], "Variable.", usage)

    display = [
        ("B1", "Set Display Font Size", "[Option]", "B2", "Sets display font size compatibility option."),
        ("B3", "Set Display Font Color", "[Foreground][Background]", "B4", "Sets display font color compatibility option."),
        ("BB", "Set Screen Saver", "[WaitTime][Type]", "BC", "Sets screen saver compatibility option."),
        ("BD", "Enable or Disable Screen Saver", "[0/1]", "BE", "Enables or disables screen saver compatibility option."),
        ("J0", "Initialize Image Table", "", "J0", "Initializes local image table."),
        ("J1", "Query Image Table", "", "J1", "Returns image table entries."),
        ("J2", "Select Images", "[Control][Names]", "J2", "Selects stored image names."),
        ("J3", "Delete Images", "[Names]", "J3", "Deletes stored images."),
        ("J4", "Download Image Packet", "[Packet]", "J4", "Stores image packet data."),
        ("J5", "Upload Image Packet", "[Control][Name]", "J5", "Returns stored image packet data."),
        ("J7", "Set Idle Image", "[Name]", "J7", "Accepted for compatibility."),
        ("J8", "Enable Idle Image", "[Option]", "J8", "Accepted for compatibility."),
        ("J9", "Show Image", "[Name]", "J9", "Shows stored image if available."),
        ("JA", "Download Boot Logo", "[Packet]", "JA", "Stores boot logo packet data."),
    ]
    for code, title, data, response, usage in display:
        add(code, title, tran, [Elem("[Data]", "Variable", data or "No payload.")], [Elem("[Status/Data]", "Variable", "Command-specific status or data.")], "Variable.", usage, response=response)
    for code, title, usage in [("BF", "Preview Screen Saver", "Compatibility no-op."), ("J6", "Play Selected Images", "Displays selected image sequence if available.")]:
        add_no_frame(code, title, tran, [Elem("[Data]", "Variable", "Command-specific payload.")], "Variable.", usage)
    prompt_cmds = [
        ("Z0", "Move Display Cursor", "[Line]", "Compatibility cursor command."),
        ("Z1", "Reset Display and Transaction State", "", "Clears prompts, transaction amount display, cached card data, active reads, and active PIN entry."),
        ("72", "Cancel Transaction", "", "Cancels active card or PIN operation and returns to idle."),
        ("Z2", "Display Prompt", "[Prompt]", "Displays one fixed or custom prompt."),
        ("Z3", "Display Multi-Line Prompt", "[Prompts]", "Displays fixed or custom prompt lines."),
        ("Z7", "Cancel Message Display Option", "[0/1]", "Controls whether cancel message is shown."),
        ("Z8", "Set Idle Prompt", "[Prompt]", "Sets the idle prompt text."),
        ("ZA", "Set Transaction Display Context", "[Type]<FS>[Amount]<FS>[Currency]<FS>[Symbol]", "Stores transaction amount, currency, and transaction type for the next card-read flow."),
        ("Z42", "Read Key Code", "[Timeout]", "Starts keypad single-key read after a prompt."),
        ("Z50", "String Entry", "[Entry control]", "Starts keypad string entry after a prompt."),
    ]
    for code, title, payload, usage in prompt_cmds:
        response = {"Z42": "Z43", "Z50": "Z51"}.get(code, code)
        add(code, title, tran, [Elem("[Data]", "Variable", payload or "No payload.")], [Elem("[Status/Data]", "Variable", "Command-specific status, key code, text, or no frame depending on command.")], "Variable.", usage, response=response, async_response=code in {"Z42", "Z50"}, final_eot=False)
    for code in ["Z43", "Z51"]:
        add(code, "Data Entry Response Identifier", tran, [Elem("[Payload]", "Variable", "Device-to-host response identifier.")], [Elem("[Payload]", "Variable", "Key code or entered text.")], "Variable.", "Device sends this after data entry completes.", final_eot=False)

    emv_setup = [
        ("T01", "Load Contact Terminal Configuration", "T02", "Stores terminal configuration data."),
        ("T03", "Load Contact CAPK", "T04", "Stores certification authority public key data."),
        ("T05", "Load Contact Application Configuration", "T06", "Stores contact application configuration."),
        ("T07", "Load EMV Data Format Table", "T08", "Stores tag format definitions for private and standard tags."),
        ("T09", "Query Contact Configuration IDs", "T0A", "Returns stored contact configuration identifiers."),
        ("T0B", "Delete Contact Configuration", "T0C", "Deletes stored contact configuration."),
        ("T51", "Load Contactless Terminal Configuration", "T52", "Stores contactless terminal configuration."),
        ("T53", "Load Contactless CAPK", "T54", "Stores contactless CAPK only when not already present."),
        ("T55", "Load Contactless Application Configuration", "T56", "Stores contactless application configuration."),
        ("T59", "Query Contactless Configuration IDs", "T5A", "Returns stored contactless configuration identifiers."),
        ("T5B", "Delete Contactless Configuration", "T5C", "Deletes stored contactless configuration."),
        ("T5D", "Contactless Housekeeping", "T5E", "Runs contactless housekeeping compatibility operation."),
        ("T5F", "Load Contactless DRL Configuration", "T5G", "Stores dynamic reader limit configuration."),
        ("T5H", "Delete Contactless DRL Configuration", "T5I", "Deletes dynamic reader limit configuration."),
    ]
    for code, title, response, usage in emv_setup:
        add(code, title, tran, [Elem("[Payload]", "Variable", "Configuration packet or command-specific selector.")], [Elem("[Status]", "1", "'0' success, '1' failure."), Elem("[Reason/Error]", "Variable", "Optional reason and error message.")], "Variable.", usage, response=response)
    emv_trans = [
        ("T11", "Contact Application Selection", "T12", "Starts contact application selection only. Multiple applications are shown to the user for touch or keypad selection."),
        ("T13", "Contact Application Selection Alias", "T12", "Alias of contact application selection."),
        ("T15", "Start Contact Transaction", "T16", "Runs contact EMV transaction. Optional PIN scheme selector may be appended."),
        ("T17", "Complete Contact Online Authorization", "T16", "Completes online authorization and second card decision."),
        ("T19", "Add Contact Issuer Script", "T20", "Stores issuer script data for completion."),
        ("T1D", "Overwrite Contact Runtime Data", "T1E", "Overwrites current transaction data objects."),
        ("T21", "Query Contact Transaction Data", "T22", "Returns requested transaction tags."),
        ("T25", "Get Contact Batch Data", "T26", "Returns next batch data record."),
        ("T27", "Get Contact Online Authorization Data", "T28", "Returns online authorization TLV data."),
        ("T29", "Get Contact Reversal Data", "T2A", "Returns reversal TLV data."),
        ("T2B", "Get Contact Online Data Packet", "T2C", "Returns online data in single final packet format."),
        ("T2H", "Get Contact Transaction Data Packet", "T2I", "Returns online, batch, or reversal data in packet format."),
        ("T31", "Contact Card Data Read", "T12", "Starts non-financial contact card data read."),
        ("T33", "Fast Contact Read", "T34", "Starts fast contact data/authorization read and returns packetized result."),
        ("T35", "Non-EMV Contact Transaction", "T16", "Starts non-EMV contact transaction flow."),
        ("T37", "Contact PIN Management", "T38", "Currently returns unsupported failure."),
        ("T3C", "Force Complete Non-EMV Contact Flow", "T3C", "Completes non-EMV card-data flow."),
        ("T61", "Start Contactless Transaction", "T62", "Runs contactless EMV transaction. Optional PIN scheme selector may be appended."),
        ("T63", "Query Contactless Transaction Data", "T64", "Returns requested contactless tags."),
        ("T65", "Query Contactless Online Data", "T66", "Returns contactless online authorization data."),
        ("T67", "Query Contactless Payment Scheme", "T68", "Returns selected payment scheme code."),
        ("T71", "Complete Contactless Online Authorization", "T62", "Completes contactless online authorization."),
        ("T73", "Add Contactless Issuer Script", "T74", "Stores contactless issuer script data."),
        ("T75", "Load Contactless Revocation List", "T76", "Loads revocation list data."),
        ("T77", "Load Contactless Exception List", "T78", "Loads exception list data."),
        ("T81", "Start Contactless EMV Transaction", "T62", "Compatibility command to start contactless EMV transaction."),
    ]
    for code, title, response, usage in emv_trans:
        add(code, title, tran, [Elem("[Payload]", "Variable", "Transaction request, host response, DOL, or command-specific data.")], [Elem("[Status]", "1", "'0' success, '1' failure."), Elem("[Result/Data]", "Variable", "Transaction result code, reversal flag, TLV data, or reason/error fields.")], "Variable.", usage, response=response, async_response=code in {"T11", "T13", "T15", "T31", "T33", "T35", "T61", "T81"}, final_eot=False)
    for code, title, usage in [("T1C", "Cancel Contact Transaction", "Cancels active contact transaction."), ("T23", "Clear Contact Transaction Log", "Clears transaction log state."), ("T34", "Fast Contact Packet Acknowledgement", "Host compatibility acknowledgement for fast contact packets."), ("T38", "PIN Management Response Identifier", "Compatibility response identifier."), ("T6C", "Cancel Contactless Transaction", "Cancels active contactless transaction."), ("T72", "Contactless Compatibility Acknowledgement", "Accepted no-op used by some hosts.")]:
        add_no_frame(code, title, tran, [], "Fixed.", usage)

    return items


def build_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.7)
    sec.right_margin = Inches(0.7)
    sec.header_distance = Inches(0.28)
    sec.footer_distance = Inches(0.28)

    styles = doc.styles
    styles["Normal"].font.name = FONT
    styles["Normal"].font.size = Pt(10.5)
    for style, size in [("Heading 1", 17), ("Heading 2", 14), ("Heading 3", 12)]:
        st = styles[style]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string("000000" if style == "Heading 1" else BLUE)
        st.paragraph_format.space_before = Pt(8)
        st.paragraph_format.space_after = Pt(6)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run(header, "NEXGO Android Smart PinPAD Programmer's Manual", size=9)
    run(header, " " * 42 + "Rev. 1.2", size=9)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(footer, "NEXGO Android Smart PinPAD Command Specification", size=8)

    p = para(doc, "NEXGO Android Smart PinPAD", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=3)
    para(doc, "Programmer's Manual", size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=16)
    para(doc, "Host Command Protocol Specification", size=13, align=WD_ALIGN_PARAGRAPH.CENTER, after=20)
    para(doc, "Revision 1.2 - 2026-07-03", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=24)
    para(doc, "This document specifies the host command protocol implemented by the NEXGO Android Smart PinPAD application. The device performs local L2 card interaction, secure PIN entry, prompts, key management, and card-reader control. Host authorization and settlement remain outside the device.", after=8)

    doc.add_page_break()
    para(doc, "Section 1    Communication Protocol", style="Heading 1")
    para(doc, "All command packets are transmitted through the selected serial transport. The application supports RS232 and USB CDC virtual serial operation. The transport layer validates start and end controls, LRC, partial-frame timeout, ACK/NAK retry, and session termination.", after=6)
    element_table(doc, "Control characters:", "Packet control values", [
        Elem("<STX>", "1", "02 hexadecimal. Starts a transaction frame."),
        Elem("<ETX>", "1", "03 hexadecimal. Ends a transaction frame."),
        Elem("<SI>", "1", "0F hexadecimal. Starts an administration frame."),
        Elem("<SO>", "1", "0E hexadecimal. Ends an administration frame."),
        Elem("<ACK>", "1", "06 hexadecimal. Positive acknowledgement."),
        Elem("<NAK>", "1", "15 hexadecimal. Transmission error acknowledgement."),
        Elem("<EOT>", "1", "04 hexadecimal. Session termination, format error, or cancellation."),
        Elem("<FS>", "1", "1C hexadecimal. Field separator."),
        Elem("<SUB>", "1", "1A hexadecimal. Subfield delimiter."),
    ])
    flow_table(doc, [
        Flow("Send request frame.", "->", ""),
        Flow("", "<-", "Send <ACK> after a valid frame is received and accepted for processing. Send <NAK> on transmission error. Send <EOT> on packet format error where specified."),
        Flow("", "<-", "Send command response frame when processing completes."),
        Flow("Send <ACK> if response is valid, <NAK> if response has LRC error, or <EOT> to cancel.", "->", ""),
        Flow("", "<-", "Administration responses send final <EOT> after host <ACK>, except commands 11 and 14. Transaction responses do not send a final <EOT>."),
    ])

    sections = [
        ("Section 2    Administration and Maintenance Messages", lambda m: m.frame == "admin"),
        ("Section 3    Security and PIN Messages", lambda m: m.code in {"20","21","22","23","24","60","62","63","70","71","7G","7A","90","91","94","96","98","99","Z60","Z62","Z64","Z65","Z66","Z67"}),
        ("Section 4    Magnetic Stripe and Multi-Interface Messages", lambda m: m.code.startswith("Q") or m.code in {"81","83"}),
        ("Section 5    CPU Smart Card Messages", lambda m: m.code.startswith("I")),
        ("Section 6    RF Memory Card Messages", lambda m: m.code.startswith("P")),
        ("Section 7    Display, Prompt, Data Entry, and Image Messages", lambda m: m.code.startswith("B") or m.code.startswith("J") or m.code.startswith("Z") or m.code == "72"),
        ("Section 8    EMV Contact and Contactless Messages", lambda m: m.code.startswith("T")),
    ]
    all_messages = messages()
    seen: set[str] = set()
    for section_title, pred in sections:
        doc.add_section(WD_SECTION.NEW_PAGE)
        para(doc, section_title, style="Heading 1")
        for msg in all_messages:
            if msg.code not in seen and pred(msg):
                seen.add(msg.code)
                message_section(doc, msg)

    doc.add_section(WD_SECTION.NEW_PAGE)
    para(doc, "Appendix A    Fixed Prompt Identifiers", style="Heading 1")
    element_table(doc, "Prompt table:", "Data entry prompts", [
        Elem("001", "3", "ACCOUNT NUMBER"),
        Elem("003", "3", "ENTER CUST ID"),
        Elem("004", "3", "ENTER AMOUNT"),
        Elem("015", "3", "ENTER CASH BACK"),
        Elem("024", "3", "ENTER"),
        Elem("100", "3", "ACCOUNT NUMBER"),
        Elem("114", "3", "CUSTOMER REF"),
        Elem("115", "3", "CUSTOMER REF NO."),
        Elem("126", "3", "ENTER BADGE #"),
        Elem("129", "3", "ENTER CASH BACK"),
        Elem("134", "3", "ENTER CUST REF"),
        Elem("200", "3", "ENTER OTP"),
        Elem("201", "3", "ENTER VERIFICATION CODE"),
        Elem("202", "3", "ENTER CVV2"),
        Elem("203", "3", "ENTER SMS OTP"),
        Elem("204", "3", "ENTER AUTHENTICATOR CODE"),
        Elem("205", "3", "ENTER TEMPORARY PASSWORD"),
    ])
    element_table(doc, "Prompt table:", "PIN prompts", [
        Elem("001", "3", "ENTER PIN"),
        Elem("002", "3", "REENTER PIN"),
        Elem("006", "3", "ENTER NEW PIN"),
        Elem("007", "3", "CONFIRM NEW PIN"),
        Elem("024", "3", "ENTER"),
    ])

    doc.core_properties.title = "NEXGO Android Smart PinPAD Programmer's Manual"
    doc.core_properties.author = "NEXGO Android Smart PinPAD"
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_doc()
